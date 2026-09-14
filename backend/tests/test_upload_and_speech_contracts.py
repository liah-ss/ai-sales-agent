import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.config import Settings
from app.core.database import Base, get_db
from app.core.security import get_current_admin_user
from app.main import app
from app.models import AdminUser, OperationLog
from app.routers import management_catalog, speech, website_config
from app.services.cos_storage import CosStorage


engine = create_engine(
    "sqlite://",
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)


def override_get_db():
    db = TestingSessionLocal()
    try:
        yield db
    finally:
        db.close()


def override_admin_user():
    return AdminUser(id=1, username="upload-admin", password_hash="test", role="super_admin", is_active=True)


client = TestClient(app)


def docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Contract document", level=1)
    document.add_paragraph("Imported body paragraph.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Parameter"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Voltage"
    table.cell(1, 1).text = "10 kV"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class UploadAndSpeechContractTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        Base.metadata.create_all(engine)

    def setUp(self) -> None:
        app.dependency_overrides[get_db] = override_get_db
        app.dependency_overrides[get_current_admin_user] = override_admin_user
        with TestingSessionLocal() as db:
            db.query(OperationLog).delete()
            db.query(AdminUser).delete()
            db.add(AdminUser(
                id=1,
                username="upload-admin",
                password_hash="test",
                role="super_admin",
                is_active=True,
            ))
            db.commit()

        self.temp_dir = tempfile.TemporaryDirectory()
        self.temp_root = Path(self.temp_dir.name)
        self.original_image_upload_dir = website_config.IMAGE_UPLOAD_DIR
        self.original_detail_upload_root = management_catalog.PRODUCT_DETAIL_UPLOAD_ROOT
        self.original_speech_upload_dir = speech.SPEECH_UPLOAD_DIR
        website_config.IMAGE_UPLOAD_DIR = self.temp_root / "website-images"
        management_catalog.PRODUCT_DETAIL_UPLOAD_ROOT = self.temp_root / "product-details"
        speech.SPEECH_UPLOAD_DIR = self.temp_root / "speech"

        local_storage = CosStorage(Settings(cos_upload_enabled=False))
        self.website_storage_patch = patch.object(website_config, "get_cos_storage", return_value=local_storage)
        self.detail_storage_patch = patch.object(management_catalog, "get_cos_storage", return_value=local_storage)
        self.website_storage_patch.start()
        self.detail_storage_patch.start()

    def tearDown(self) -> None:
        self.website_storage_patch.stop()
        self.detail_storage_patch.stop()
        website_config.IMAGE_UPLOAD_DIR = self.original_image_upload_dir
        management_catalog.PRODUCT_DETAIL_UPLOAD_ROOT = self.original_detail_upload_root
        speech.SPEECH_UPLOAD_DIR = self.original_speech_upload_dir
        self.temp_dir.cleanup()

    def test_website_image_and_docx_upload_contracts(self) -> None:
        image = client.post(
            "/api/management/website-config/uploads/images",
            files={"file": ("contract.png", b"png-contract", "image/png")},
        )
        self.assertEqual(image.status_code, 200)
        self.assertEqual(image.json()["fileName"], "contract.png")
        self.assertTrue(image.json()["url"].startswith("/uploads/images/"))
        self.assertEqual(len(list((self.temp_root / "website-images").glob("*.png"))), 1)

        parsed = client.post(
            "/api/management/website-config/uploads/docx/parse",
            files={
                "file": (
                    "contract.docx",
                    docx_bytes(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        self.assertEqual(parsed.status_code, 200)
        self.assertEqual(parsed.json()["title"], "Contract document")
        self.assertEqual(parsed.json()["tables"][0][1], ["Voltage", "10 kV"])
        self.assertEqual([block["type"] for block in parsed.json()["blocks"]], ["paragraph", "specs"])

    def test_product_detail_asset_upload_contract(self) -> None:
        response = client.post(
            "/api/management/catalog/products/detail-assets",
            files={"file": ("detail.png", b"product-detail", "image/png")},
        )
        self.assertEqual(response.status_code, 201)
        block = response.json()["blocks"][0]
        self.assertEqual(block["type"], "image")
        self.assertEqual(block["name"], "detail.png")
        self.assertTrue(block["url"].startswith("/uploads/product-details/"))
        self.assertEqual(len(list((self.temp_root / "product-details").rglob("*.png"))), 1)

    def test_speech_transcription_success_contract(self) -> None:
        with patch.object(speech, "call_transcription_api", return_value="recognized contract speech") as provider:
            response = client.post(
                "/api/speech/transcriptions",
                files={"audio": ("speech.webm", b"audio-contract", "audio/webm")},
                data={"language": "zh"},
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {"text": "recognized contract speech"})
        provider.assert_called_once_with(b"audio-contract", "speech.webm", "audio/webm", "zh")
        self.assertEqual(len(list((self.temp_root / "speech").glob("*-speech.webm"))), 1)

    def test_realtime_speech_websocket_contract(self) -> None:
        class FakeRecognition:
            frames: list[bytes] = []
            started = False
            stopped = False

            def __init__(self, **_kwargs) -> None:
                pass

            def start(self) -> None:
                FakeRecognition.started = True

            def send_audio_frame(self, content: bytes) -> None:
                FakeRecognition.frames.append(content)

            def stop(self) -> None:
                FakeRecognition.stopped = True

        settings = SimpleNamespace(
            dashscope_api_key="contract-key",
            dashscope_websocket_url="wss://example.invalid/asr",
            realtime_asr_model="contract-asr",
        )
        with (
            patch.object(speech, "get_settings", return_value=settings),
            patch.object(speech, "Recognition", FakeRecognition),
            client.websocket_connect("/api/speech/realtime") as websocket,
        ):
            self.assertEqual(websocket.receive_json(), {"type": "ready"})
            websocket.send_bytes(b"pcm-frame")
            websocket.send_text("stop")

        self.assertTrue(FakeRecognition.started)
        self.assertEqual(FakeRecognition.frames, [b"pcm-frame"])
        self.assertTrue(FakeRecognition.stopped)


if __name__ == "__main__":
    unittest.main()
