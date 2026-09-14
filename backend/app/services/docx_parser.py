from io import BytesIO

from docx import Document


def parse_docx_bytes(content: bytes, file_name: str) -> dict[str, object]:
    document = Document(BytesIO(content))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    tables: list[list[list[str]]] = []

    for table in document.tables:
        rows: list[list[str]] = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        if rows:
            tables.append(rows)

    title = paragraphs[0] if paragraphs else file_name.rsplit(".", 1)[0]
    body = "\n\n".join(paragraphs[1:] if len(paragraphs) > 1 else paragraphs)
    blocks = [
        {
            "id": f"docx-{abs(hash(file_name))}-overview",
            "type": "paragraph",
            "title": title,
            "body": body,
        }
    ]

    if tables:
        blocks.append(
            {
                "id": f"docx-{abs(hash(file_name))}-table",
                "type": "specs",
                "title": "Imported table",
                "body": "\n".join(" | ".join(row) for row in tables[0]),
            }
        )

    return {
        "fileName": file_name,
        "title": title,
        "paragraphs": paragraphs,
        "tables": tables,
        "blocks": blocks,
    }
